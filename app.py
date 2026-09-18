import os
import re
import threading
import subprocess
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import pymupdf
from docx import Document
from rapidfuzz import fuzz


ARABIC_DIACRITICS = re.compile(r"[\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED]")
WHITESPACE = re.compile(r"\s+")

def normalize_arabic(text: str) -> str:
    if not text:
        return ""
    text = ARABIC_DIACRITICS.sub("", text)
    text = text.replace("ـ", "")
    text = (text
            .replace("أ", "ا")
            .replace("إ", "ا")
            .replace("آ", "ا")
            .replace("ٱ", "ا")
            .replace("ى", "ي")
            .replace("ة", "ه"))
    # Normalize Arabic-Indic digits and Persian digits to ASCII.
    trans = str.maketrans(
        "٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹",
        "01234567890123456789"
    )
    text = text.translate(trans)
    text = WHITESPACE.sub(" ", text)
    return text.strip().casefold()

def compact(text: str) -> str:
    return re.sub(r"\s+", "", normalize_arabic(text))

def context(text: str, start: int, end: int, radius=260) -> str:
    a = max(0, start - radius)
    b = min(len(text), end + radius)
    return text[a:b].replace("\n", " ").strip()

def extract_docx(path: Path):
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    # Include table cells because legal documents may put case data in tables.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return [(1, "\n".join(parts))]

def extract_pdf(path: Path, use_ocr=False, tesseract_path=""):
    pages = []
    doc = pymupdf.open(str(path))
    try:
        if tesseract_path:
            # PyMuPDF uses Tesseract through its OCR support.
            os.environ["PATH"] = str(Path(tesseract_path).parent) + os.pathsep + os.environ.get("PATH", "")
        for i, page in enumerate(doc):
            text = page.get_text("text", sort=True)
            if use_ocr and len(text.strip()) < 20:
                try:
                    tp = page.get_textpage_ocr(language="ara+eng", dpi=200, full=True)
                    text = page.get_text("text", textpage=tp, sort=True)
                except Exception:
                    # Keep extracted text if OCR is unavailable.
                    pass
            pages.append((i + 1, text))
    finally:
        doc.close()
    return pages

def extract(path: Path, use_ocr=False, tesseract_path=""):
    if path.suffix.lower() == ".pdf":
        return extract_pdf(path, use_ocr, tesseract_path)
    if path.suffix.lower() == ".docx":
        return extract_docx(path)
    return []

class SearchApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Arabic PDF + DOCX Search")
        self.geometry("1180x760")
        self.minsize(900, 600)

        self.folder = tk.StringVar()
        self.query = tk.StringVar()
        self.mode = tk.StringVar(value="normalized")
        self.use_ocr = tk.BooleanVar(value=False)
        self.tesseract = tk.StringVar()
        self.status = tk.StringVar(value="Choose a folder and index your documents.")
        self.index = []  # (path, page, original_text, normalized_text)
        self.excluded_folders = []
        self.index_button = None
        self.progress = None

        self.build_ui()

    def build_ui(self):
        pad = {"padx": 10, "pady": 6}

        top = ttk.Frame(self)
        top.pack(fill="x", **pad)

        ttk.Label(top, text="Documents folder:").pack(side="left")
        ttk.Entry(top, textvariable=self.folder).pack(side="left", fill="x", expand=True, padx=8)
        ttk.Button(top, text="Choose Folder", command=self.choose_folder).pack(side="left")
        ttk.Button(top, text="Exclude Folders", command=self.manage_excluded_folders).pack(side="left", padx=5)
        self.index_button = ttk.Button(top, text="Index Files", command=self.start_index)
        self.index_button.pack(side="left", padx=5)

        opts = ttk.Frame(self)
        opts.pack(fill="x", **pad)

        ttk.Label(opts, text="Search:").pack(side="left")
        q = ttk.Entry(opts, textvariable=self.query, font=("Segoe UI", 12))
        q.pack(side="left", fill="x", expand=True, padx=8)
        q.bind("<Return>", lambda e: self.search())

        ttk.Button(opts, text="Search", command=self.search).pack(side="left")

        ttk.Label(opts, text="Mode:").pack(side="left", padx=(12, 4))
        ttk.Combobox(
            opts, textvariable=self.mode, state="readonly", width=18,
            values=("normalized", "fuzzy")
        ).pack(side="left")

        ocr = ttk.Frame(self)
        ocr.pack(fill="x", **pad)
        ttk.Checkbutton(
            ocr, text="OCR scanned PDFs when extracted text is missing",
            variable=self.use_ocr
        ).pack(side="left")
        ttk.Label(ocr, text="Tesseract path (optional):").pack(side="left", padx=(20, 4))
        ttk.Entry(ocr, textvariable=self.tesseract, width=45).pack(side="left")
        ttk.Button(ocr, text="Browse", command=self.choose_tesseract).pack(side="left", padx=5)

        ttk.Label(self, textvariable=self.status).pack(fill="x", **pad)

        self.progress = ttk.Progressbar(self, mode="determinate", maximum=1, value=0)
        self.progress.pack(fill="x", padx=10, pady=(0, 6))

        columns = ("file", "page", "score", "context")
        self.tree = ttk.Treeview(self, columns=columns, show="headings")
        self.tree.heading("file", text="File")
        self.tree.heading("page", text="Page")
        self.tree.heading("score", text="Match")
        self.tree.heading("context", text="Context")
        self.tree.column("file", width=280)
        self.tree.column("page", width=70, anchor="center")
        self.tree.column("score", width=80, anchor="center")
        self.tree.column("context", width=680)
        self.tree.pack(fill="both", expand=True, padx=10, pady=(4, 0))
        self.tree.bind("<Double-1>", self.open_result)

        scroll = ttk.Scrollbar(self, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        scroll.place(relx=0.985, rely=0.28, relheight=0.62)

    def choose_folder(self):
        p = filedialog.askdirectory(title="Choose folder containing PDFs/DOCX")
        if p:
            self.folder.set(p)

    def manage_excluded_folders(self):
        root_folder = self.folder.get().strip()
        dialog = tk.Toplevel(self)
        dialog.title("Excluded folders")
        dialog.transient(self)
        dialog.grab_set()
        dialog.geometry("620x300")

        ttk.Label(
            dialog,
            text="Folders below the selected documents folder will be skipped during indexing.",
            wraplength=580,
        ).pack(fill="x", padx=12, pady=(12, 6))

        listbox = tk.Listbox(dialog, height=8, exportselection=False)
        listbox.pack(fill="both", expand=True, padx=12, pady=6)
        for excluded in self.excluded_folders:
            listbox.insert("end", excluded)

        controls = ttk.Frame(dialog)
        controls.pack(fill="x", padx=12, pady=(0, 12))

        def add_folder():
            selected = filedialog.askdirectory(
                parent=dialog,
                title="Choose folder to exclude",
                initialdir=root_folder or None,
            )
            if not selected:
                return
            selected_path = Path(selected).resolve()
            if root_folder:
                root_path = Path(root_folder).resolve()
                try:
                    selected_path.relative_to(root_path)
                except ValueError:
                    messagebox.showwarning(
                        "Invalid exclusion",
                        "Excluded folders must be inside the selected documents folder.",
                        parent=dialog,
                    )
                    return
            selected_text = str(selected_path)
            if selected_text not in [listbox.get(i) for i in range(listbox.size())]:
                listbox.insert("end", selected_text)

        def remove_folder():
            selection = listbox.curselection()
            if selection:
                listbox.delete(selection[0])

        ttk.Button(controls, text="Add Folder", command=add_folder).pack(side="left")
        ttk.Button(controls, text="Remove Selected", command=remove_folder).pack(side="left", padx=6)

        def save():
            self.excluded_folders = list(listbox.get(0, "end"))
            dialog.destroy()

        ttk.Button(controls, text="Done", command=save).pack(side="right")

    def choose_tesseract(self):
        p = filedialog.askopenfilename(
            title="Select tesseract.exe",
            filetypes=[("Tesseract", "tesseract.exe"), ("All files", "*.*")]
        )
        if p:
            self.tesseract.set(p)

    def start_index(self):
        folder_value = self.folder.get().strip()
        if not folder_value:
            messagebox.showwarning("Folder required", "Choose a documents folder first.")
            return

        self.tree.delete(*self.tree.get_children())
        self.index = []
        self.index_button.configure(state="disabled")
        self.progress.configure(mode="indeterminate", value=0)
        self.progress.start(12)

        use_ocr = self.use_ocr.get()
        tesseract_path = self.tesseract.get().strip()
        excluded_folders = tuple(self.excluded_folders)
        threading.Thread(
            target=self.index_documents,
            args=(Path(folder_value), use_ocr, tesseract_path, excluded_folders),
            daemon=True,
        ).start()

    def index_documents(self, folder, use_ocr, tesseract_path, excluded_folders=()):
        try:
            excluded = {Path(path).resolve() for path in excluded_folders}
            paths = []
            for current_root, directories, filenames in os.walk(folder):
                current_path = Path(current_root).resolve()
                directories[:] = [
                    name for name in directories
                    if (current_path / name).resolve() not in excluded
                ]
                paths.extend(
                    current_path / name
                    for name in filenames
                    if Path(name).suffix.lower() in {".pdf", ".docx"}
                )
        except Exception as e:
            self.after(0, self.finish_index, [], 0, f"Could not scan folder: {e}")
            return

        total = len(paths)
        self.after(0, self.prepare_progress, total)
        new_index = []

        for n, path in enumerate(paths, 1):
            try:
                pages = extract(path, use_ocr=use_ocr, tesseract_path=tesseract_path)
                for page_no, text in pages:
                    if text.strip():
                        new_index.append((str(path), page_no, text, normalize_arabic(text)))
            except Exception as e:
                print(f"ERROR: {path}: {e}")
            self.after(0, self.update_progress, n, total, path.name)

        self.after(
            0,
            self.finish_index,
            new_index,
            total,
            f"Indexed {len(new_index)} searchable pages/sections from {total} files.",
        )

    def prepare_progress(self, total):
        self.progress.stop()
        self.progress.configure(mode="determinate", maximum=max(1, total), value=0)
        self.status.set(f"Found {total} PDF/DOCX files. Starting index...")

    def update_progress(self, current, total, filename):
        self.progress.configure(value=current)
        self.status.set(f"Indexing {current}/{total}: {filename}")

    def finish_index(self, new_index, total, status):
        self.index = new_index
        self.progress.stop()
        self.progress.configure(mode="determinate", maximum=max(1, total), value=total)
        self.status.set(status)
        self.index_button.configure(state="normal")

    def search(self):
        q = self.query.get().strip()
        if not q:
            return
        if not self.index:
            messagebox.showinfo("No index", "Index the document folder first.")
            return

        nq = normalize_arabic(q)
        cq = compact(q)
        mode = self.mode.get()
        results = []

        for path, page, original, norm in self.index:
            score = 0
            pos = -1

            if nq in norm:
                score = 100
                pos = norm.find(nq)
            elif cq and cq in compact(original):
                score = 99
                pos = 0
            elif mode == "fuzzy":
                # Compare against sliding word windows so long legal phrases work.
                words = norm.split()
                qwords = nq.split()
                if qwords and words:
                    window = len(qwords)
                    best = 0
                    best_pos = 0
                    for i in range(max(1, len(words) - window + 1)):
                        candidate = " ".join(words[i:i + window + 4])
                        s = fuzz.partial_ratio(nq, candidate)
                        if s > best:
                            best = s
                            best_pos = i
                    if best >= 72:
                        score = best
                        # Context is generated from original text below.
                        pos = 0

            if score:
                # Map approximate normalized position back to original text.
                if pos >= 0 and pos < len(norm):
                    prefix = norm[:pos]
                    ratio = len(prefix) / max(1, len(norm))
                    orig_pos = int(ratio * len(original))
                else:
                    orig_pos = 0
                results.append((score, path, page, context(original, orig_pos, orig_pos + len(q))))

        results.sort(key=lambda x: (-x[0], x[1], x[2]))
        self.tree.delete(*self.tree.get_children())

        for score, path, page, ctx in results[:500]:
            self.tree.insert("", "end", values=(path, page, f"{score:.0f}%", ctx))

        self.status.set(f"Found {len(results)} matching pages/sections.")

    def open_result(self, _event=None):
        item = self.tree.focus()
        if not item:
            return
        path = self.tree.item(item, "values")[0]
        try:
            os.startfile(path)
        except Exception:
            subprocess.Popen(["explorer", "/select,", path])

if __name__ == "__main__":
    SearchApp().mainloop()
